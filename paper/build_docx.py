"""Build the editable Word manuscript from the maintained LaTeX source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscript.tex"
OUTPUT = ROOT / "docx" / "QWeave_Manuscript.docx"
LEARNING_CURVE = ROOT.parent / "results" / "benchmark_v2" / "figures" / "learning_curves.png"

CITATIONS = {
    "childs2019transformations": "Childs et al., 2019",
    "coppersmith2002qft": "Coppersmith, 2002",
    "cowtan2019routing": "Cowtan et al., 2019",
    "farhi2014qaoa": "Farhi et al., 2014",
    "ibmsabre": "IBM Quantum, 2026a",
    "ibmtranspiler": "IBM Quantum, 2026b",
    "kipf2017gcn": "Kipf and Welling, 2017",
    "li2019sabre": "Li et al., 2019",
    "nannicini2022ip": "Nannicini et al., 2022",
    "pozzi2020rl": "Pozzi et al., 2020",
    "schulman2015gae": "Schulman et al., 2015",
    "schulman2017ppo": "Schulman et al., 2017",
    "sinha2022gnn": "Sinha et al., 2022",
    "siraichi2018allocation": "Siraichi et al., 2018",
    "zulehner2019efficient": "Zulehner et al., 2019",
}
REFERENCES = [
    "Childs, A. M., Schoute, E., and Unsal, C. M. (2019). Circuit Transformations for Quantum Architectures. TQC 2019, LIPIcs 135, 3:1-3:24. doi:10.4230/LIPIcs.TQC.2019.3.",
    "Coppersmith, D. (2002). An Approximate Fourier Transform Useful in Quantum Factoring. arXiv:quant-ph/0201067.",
    "Cowtan, A., Dilkes, S., Duncan, R., Krajenbrink, A., Simmons, W., and Sivarajah, S. (2019). On the Qubit Routing Problem. TQC 2019, LIPIcs 135, 5:1-5:32. doi:10.4230/LIPIcs.TQC.2019.5.",
    "Farhi, E., Goldstone, J., and Gutmann, S. (2014). A Quantum Approximate Optimization Algorithm. arXiv:1411.4028.",
    "IBM Quantum. (2026a). SabreLayout API Reference. IBM Quantum Documentation.",
    "IBM Quantum. (2026b). Transpiler Stages. IBM Quantum Documentation.",
    "Kipf, T. N., and Welling, M. (2017). Semi-Supervised Classification with Graph Convolutional Networks. ICLR 2017. arXiv:1609.02907.",
    "Li, G., Ding, Y., and Xie, Y. (2019). Tackling the Qubit Mapping Problem for NISQ-Era Quantum Devices. ASPLOS 2019. arXiv:1809.02573.",
    "Nannicini, G., Bishop, L. S., Gunluk, O., and Jurcevic, P. (2022). Optimal Qubit Assignment and Routing via Integer Programming. ACM Transactions on Quantum Computing, 4(1). doi:10.1145/3544563.",
    "Pozzi, M. G., Herbert, S. J., Sengupta, A., and Mullins, R. D. (2020). Using Reinforcement Learning to Perform Qubit Routing in Quantum Compilers. arXiv:2007.15957.",
    "Schulman, J., Moritz, P., Levine, S., Jordan, M. I., and Abbeel, P. (2015). High-Dimensional Continuous Control Using Generalized Advantage Estimation. arXiv:1506.02438.",
    "Schulman, J., Wolski, F., Dhariwal, P., Radford, A., and Klimov, O. (2017). Proximal Policy Optimization Algorithms. arXiv:1707.06347.",
    "Sinha, A., Azad, U., and Singh, H. (2022). Qubit Routing Using Graph Neural Network Aided Monte Carlo Tree Search. AAAI 2022. arXiv:2104.01992.",
    "Siraichi, M. Y., dos Santos, V. F., Collange, S., and Pereira, F. M. Q. (2018). Qubit Allocation. CGO 2018, 113-125. doi:10.1145/3168822.",
    "Zulehner, A., Paler, A., and Wille, R. (2019). An Efficient Methodology for Mapping Quantum Circuits to the IBM QX Architectures. IEEE TCAD, 38(7), 1226-1236. doi:10.1109/TCAD.2018.2846658.",
]
AUTHORS = (
    ("Atharva Sheersh Pandey", "1"),
    ("Diptesh Das", "1"),
    ("Shrivardhini N", "1"),
    ("Haridasu Sreedhar", "1"),
    ("Prof. Bhuvaneswari M", "2"),
)
STUDENT_EMAILS = (
    "atharva.sheersh2024@vitstudent.ac.in",
    "diptesh.das2024@vitstudent.ac.in",
    "shrivardhini.n2024@vitstudent.ac.in",
    "haridasu.sreedhar2024@vitstudent.ac.in",
)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, text, end))


def _set_columns(section, count: int) -> None:
    section_properties = section._sectPr
    columns = section_properties.xpath("./w:cols")
    element = columns[0] if columns else OxmlElement("w:cols")
    element.set(qn("w:num"), str(count))
    element.set(qn("w:space"), "340")
    if not columns:
        section_properties.append(element)


def _clean(text: str) -> str:
    def citation(match) -> str:
        command, keys = match.groups()
        labels = [CITATIONS[key.strip()] for key in keys.split(",")]
        if command == "citet" and len(labels) == 1:
            author, year = labels[0].rsplit(", ", 1)
            return f"{author} ({year})"
        return "(" + "; ".join(labels) + ")"

    replacements = {
        r"Eq.~\eqref{eq:objective}": "Equation (1)",
        r"Equation~\eqref{eq:route-invariant}": "Equation (2)",
        r"Equation~\eqref{eq:reverse-cx}": "Equation (4)",
        r"Table~\ref{tab:test-results}": "Table 1",
        r"Figure~\ref{fig:learning-curves}": "Figure 1",
        r"\rightarrow": "->", r"\leq": "<=", r"\ldots": "...", r"\emph": "",
        r"\texttt": "", r"\textsc": "", r"\small": "",
        r"\mathrm": "", r"\circ": " o ", r"\otimes": " x ",
        r"\{": "{", r"\}": "}", r"\%": "%", r"\_": "_", r"~": " ",
    }
    text = re.sub(r"\\(citep|citet)\{([^}]+)\}", citation, text)
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\(?:label|ref)\{[^}]+\}", "", text)
    text = re.sub(r"\\[A-Za-z]+", "", text)
    text = (text.replace("{", "").replace("}", "").replace("$", "")
            .replace("--", "-"))
    return " ".join(text.split())


def _add_body_paragraph(
    document: Document,
    text: str,
    *,
    first_after_heading: bool = False,
    font_size: float = 11,
    line_spacing: float = 13.6,
    left_indent: float = 0,
    right_indent: float = 0,
) -> None:
    cleaned = _clean(text)
    if not cleaned:
        return
    paragraph = document.add_paragraph(cleaned)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0 if first_after_heading else 0.38)
    paragraph.paragraph_format.left_indent = Cm(left_indent)
    paragraph.paragraph_format.right_indent = Cm(right_indent)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(line_spacing)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0, 0, 0)


def _add_author_block(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    for index, (name, marker) in enumerate(AUTHORS):
        if index:
            paragraph.add_run(", ")
        if index == 3:
            paragraph.add_run().add_break()
        name_run = paragraph.add_run(name)
        name_run.font.name = "Times New Roman"
        name_run.font.size = Pt(12)
        name_run.font.bold = True
        name_run.font.color.rgb = RGBColor(0, 0, 0)
        marker_run = paragraph.add_run(marker)
        marker_run.font.name = "Times New Roman"
        marker_run.font.size = Pt(8)
        marker_run.font.superscript = True
        marker_run.font.color.rgb = RGBColor(0, 0, 0)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    for index, email in enumerate(STUDENT_EMAILS):
        if index:
            paragraph.add_run("; ")
        if index == 2:
            paragraph.add_run().add_break()
        run = paragraph.add_run(email)
        run.font.name = "Courier New"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)

    for marker, text in (
        ("1", "School of Computer Science and Engineering, VIT Vellore"),
        ("2", "Faculty Guide, School of Computer Science and Engineering, VIT Vellore"),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        marker_run = paragraph.add_run(marker)
        marker_run.font.name = "Times New Roman"
        marker_run.font.size = Pt(8)
        marker_run.font.superscript = True
        marker_run.font.color.rgb = RGBColor(0, 0, 0)
        text_run = paragraph.add_run(text)
        text_run.font.name = "Times New Roman"
        text_run.font.size = Pt(12)
        text_run.font.color.rgb = RGBColor(0, 0, 0)


def _set_table_horizontal_rules(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for name in ("top", "insideH", "bottom"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), "000000")
        borders.append(edge)
    for name in ("left", "insideV", "right"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)


def _add_results_table(document: Document) -> None:
    wide = document.add_section(WD_SECTION.CONTINUOUS)
    wide.page_width, wide.page_height = Cm(21), Cm(29.7)
    wide.top_margin = wide.bottom_margin = Cm(2.54)
    wide.left_margin = wide.right_margin = Cm(2.54)
    _set_columns(wide, 1)
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_horizontal_rules(table)
    headers = ("Method", "Depth", "SWAP", "Runtime", "Fallback")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0, 0, 0)
    rows = (
        ("Basic", "48.5", "11.5", "0.0022", "0"),
        ("Weighted", "36.0", "0.0", "0.0036", "0"),
        ("SABRE", "46.0", "7.0", "0.0028", "0"),
        ("GNN-PPO", "37.5", "0.0", "0.3473", "4543"),
        ("No-message PPO", "36.0", "0.0", "0.1040", "742"),
        ("Untrained GNN", "198.0", "193.0", "0.6681", "8059"),
    )
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.LEFT
                                            if cell is cells[0]
                                            else WD_ALIGN_PARAGRAPH.CENTER)
            for run in cell.paragraphs[0].runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0, 0, 0)
    caption = document.add_paragraph(
        "Table 1  Benchmark-v2 test summary. Runtime is median wall time in seconds."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(0)
    caption.runs[0].font.name = "Times New Roman"
    caption.runs[0].font.size = Pt(10)
    caption.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    columns = document.add_section(WD_SECTION.CONTINUOUS)
    columns.page_width, columns.page_height = Cm(21), Cm(29.7)
    columns.top_margin = columns.bottom_margin = Cm(2.54)
    columns.left_margin = columns.right_margin = Cm(2.54)
    _set_columns(columns, 2)


def _add_learning_figure(document: Document) -> None:
    if not LEARNING_CURVE.exists():
        raise FileNotFoundError(f"missing benchmark figure: {LEARNING_CURVE}")
    wide = document.add_section(WD_SECTION.CONTINUOUS)
    wide.page_width, wide.page_height = Cm(21), Cm(29.7)
    wide.top_margin = wide.bottom_margin = Cm(2.54)
    wide.left_margin = wide.right_margin = Cm(2.54)
    _set_columns(wide, 1)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(LEARNING_CURVE), width=Cm(15.6))
    caption = document.add_paragraph(
        "Figure 1  Five-seed learning curves. Lines show the median recent "
        "episode return and bands show the interquartile range.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].font.name = "Times New Roman"
    caption.runs[0].font.size = Pt(10)
    caption.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    columns = document.add_section(WD_SECTION.CONTINUOUS)
    columns.page_width, columns.page_height = Cm(21), Cm(29.7)
    columns.top_margin = columns.bottom_margin = Cm(2.54)
    columns.left_margin = columns.right_margin = Cm(2.54)
    _set_columns(columns, 2)


def build() -> Path:
    source = SOURCE.read_text(encoding="utf-8")
    title = re.search(r"\\title\{(.+?)\}", source).group(1)
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", source, re.S).group(1)
    body = source.split(r"\end{abstract}", 1)[1].split(r"\bibliographystyle", 1)[0]
    body = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", "\n\n[[RESULTS_TABLE]]\n\n", body, flags=re.S)
    body = re.sub(r"\\begin\{figure\*\}.*?\\end\{figure\*\}",
                  "\n\n[[LEARNING_FIGURE]]\n\n", body, flags=re.S)
    equations = [
        "J(m) = sum over (i,j) in E_C of w_ij d_H(m(i),m(j)).   (1)",
        "U_(R_k) P_(m0) = P_(mk) U_(C_k), for 0 <= k <= T.   (2)",
        "r_i = d_i + max over successors j of r_j;   s_i = max over predecessors j of (s_j + d_j).   (3)",
        "CX_(u -> v) = (H_u x H_v) CX_(v -> u) (H_u x H_v).   (4)",
    ]
    for equation in equations:
        body = re.sub(r"\\begin\{equation\}.*?\\end\{equation\}",
                      f"\n\n[[EQUATION:{equation}]]\n\n", body, count=1, flags=re.S)

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = Pt(0)
    styles["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["Normal"].paragraph_format.line_spacing = Pt(13.6)
    for name, size in (("Title", 14.5), ("Heading 1", 12), ("Heading 2", 11)):
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = paragraph.add_run(_clean(title))
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14.5)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.paragraph_format.space_after = Pt(5)
    _add_author_block(document)

    columns = document.add_section(WD_SECTION.CONTINUOUS)
    columns.page_width, columns.page_height = Cm(21), Cm(29.7)
    columns.top_margin = columns.bottom_margin = Cm(2.54)
    columns.left_margin = columns.right_margin = Cm(2.54)
    _set_columns(columns, 2)

    heading = document.add_paragraph("Abstract")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.name = "Times New Roman"
    heading.runs[0].font.size = Pt(12)
    heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    _add_body_paragraph(
        document, abstract, first_after_heading=True,
        font_size=10, line_spacing=12, left_indent=0.6, right_indent=0.6,
    )

    tokens = re.split(
        r"(\\section\{[^}]+\}|\\subsection\{[^}]+\}|"
        r"\\paragraph\{[^}]+\}|\[\[RESULTS_TABLE\]\]|"
        r"\[\[LEARNING_FIGURE\]\]|\[\[EQUATION:.*?\]\])",
        body,
    )
    section_number = 0
    subsection_number = 0
    first_after_heading = False
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        section_match = re.fullmatch(r"\\section\{([^}]+)\}", token)
        subsection_match = re.fullmatch(r"\\subsection\{([^}]+)\}", token)
        paragraph_match = re.fullmatch(r"\\paragraph\{([^}]+)\}", token)
        if section_match:
            section_number += 1
            subsection_number = 0
            document.add_heading(
                f"{section_number} {_clean(section_match.group(1))}", level=1
            )
            first_after_heading = True
        elif subsection_match:
            subsection_number += 1
            document.add_heading(
                f"{section_number}.{subsection_number} "
                f"{_clean(subsection_match.group(1))}",
                level=2,
            )
            first_after_heading = True
        elif paragraph_match:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(_clean(paragraph_match.group(1)))
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
            first_after_heading = True
        elif token == "[[RESULTS_TABLE]]":
            _add_results_table(document)
            first_after_heading = True
        elif token == "[[LEARNING_FIGURE]]":
            _add_learning_figure(document)
            first_after_heading = True
        elif token.startswith("[[EQUATION:"):
            equation = token[len("[[EQUATION:"):-2]
            paragraph = document.add_paragraph(equation)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.runs[0].font.name = "Cambria Math"
            paragraph.runs[0].font.size = Pt(11)
            paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)
            first_after_heading = False
        else:
            for block in re.split(r"\n\s*\n", token):
                if _clean(block):
                    _add_body_paragraph(
                        document, block, first_after_heading=first_after_heading
                    )
                    first_after_heading = False

    document.add_heading("References", level=1)
    for reference in REFERENCES:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.35)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(12)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
    document.core_properties.title = _clean(title)
    document.core_properties.author = "; ".join(name for name, _ in AUTHORS)
    document.core_properties.subject = "Correctness-first qubit mapping and routing study"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
