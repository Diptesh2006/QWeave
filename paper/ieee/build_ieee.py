"""Build IEEE conference-style DOCX, PDF, and IEEEtran source for QWeave."""

from __future__ import annotations

from html import escape
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    FrameBreak,
    Image,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Table,
    TableStyle,
)
from PIL import Image as PILImage


HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
REPO = PAPER.parent
SOURCE = PAPER / "manuscript.tex"
FIGURE = REPO / "results" / "benchmark_v2" / "figures" / "learning_curves.png"
DOCX_OUTPUT = HERE / "docx" / "QWeave_IEEE_Conference.docx"
PDF_OUTPUT = HERE / "pdf" / "QWeave_IEEE_Conference.pdf"
TEX_OUTPUT = HERE / "latex" / "QWeave_IEEE_Conference.tex"

TITLE = "QWeave: Correctness-First Deterministic and Learned Qubit Routing"
KEYWORDS = (
    "quantum compilation, qubit mapping, qubit routing, reinforcement learning, "
    "graph neural networks, scheduling"
)
AUTHORS = (
    ("Atharva Sheersh Pandey", "1"),
    ("Diptesh Das", "1"),
    ("Shrivardhini N", "1"),
    ("Haridasu Sreedhar", "1"),
    ("Prof. Bhuvaneswari M", "2"),
)
EMAILS = (
    "atharva.sheersh2024@vitstudent.ac.in",
    "diptesh.das2024@vitstudent.ac.in",
    "shrivardhini.n2024@vitstudent.ac.in",
    "haridasu.sreedhar2024@vitstudent.ac.in",
)

# IEEE references follow first-citation order in the manuscript.
REFERENCE_KEYS = (
    "ibmtranspiler",
    "li2019sabre",
    "pozzi2020rl",
    "sinha2022gnn",
    "preskill2018nisq",
    "bharti2022nisq",
    "nielsenchuang2010",
    "javadiabhari2024qiskit",
    "khammassi2021openql",
    "sivarajah2020tket",
    "bandic2022fullstack",
    "siraichi2018allocation",
    "zulehner2019efficient",
    "cowtan2019routing",
    "childs2019transformations",
    "wille2016lookahead",
    "lye2015minimal",
    "li2021subgraph",
    "itoko2020optimization",
    "tan2021gateabsorption",
    "steinberg2022topological",
    "ibmsabre",
    "murali2019noise",
    "tannu2019variability",
    "nishio2020erroraware",
    "lao2022timing",
    "guerreschi2018twostep",
    "murali2019fullstack",
    "hillmich2021teleportation",
    "baker2020partitioning",
    "bandic2023qubo",
    "laobrowne2021twoqan",
    "bandic2020dse",
    "herbert2018rl",
    "acampora2021neural",
    "kipf2017gcn",
    "gilmer2017mpnn",
    "schulman2017ppo",
    "schulman2015gae",
    "nannicini2022ip",
    "venturelli2018temporal",
    "venturelli2019reasoning",
    "bandic2023characterization",
    "cross2019volume",
    "blumekohout2020volumetric",
    "mills2021holistic",
    "tomesh2022supermarq",
    "li2020qasmbench",
    "farhi2014qaoa",
    "coppersmith2002qft",
)
CITE_NUM = {key: index + 1 for index, key in enumerate(REFERENCE_KEYS)}
REFERENCES = (
    "IBM Quantum, \"Transpiler stages,\" IBM Quantum Documentation, 2026.",
    "G. Li, Y. Ding, and Y. Xie, \"Tackling the qubit mapping problem for NISQ-era quantum devices,\" in Proc. ASPLOS, 2019.",
    "M. G. Pozzi, S. J. Herbert, A. Sengupta, and R. D. Mullins, \"Using reinforcement learning to perform qubit routing in quantum compilers,\" arXiv:2007.15957, 2020.",
    "A. Sinha, U. Azad, and H. Singh, \"Qubit routing using graph neural network aided Monte Carlo tree search,\" in Proc. AAAI, 2022.",
    "J. Preskill, \"Quantum computing in the NISQ era and beyond,\" Quantum, vol. 2, p. 79, 2018.",
    "K. Bharti et al., \"Noisy intermediate-scale quantum algorithms,\" Rev. Mod. Phys., vol. 94, no. 1, p. 015004, 2022.",
    "M. A. Nielsen and I. L. Chuang, Quantum Computation and Quantum Information, 10th anniversary ed. Cambridge, U.K.: Cambridge Univ. Press, 2010.",
    "A. Javadi-Abhari et al., \"Quantum computing with Qiskit,\" arXiv:2405.08810, 2024.",
    "N. Khammassi et al., \"OpenQL: A portable quantum programming framework for quantum accelerators,\" ACM J. Emerg. Technol. Comput. Syst., vol. 18, no. 1, pp. 1-24, 2021.",
    "S. Sivarajah, S. Dilkes, A. Cowtan, W. Simmons, A. Edgington, and R. Duncan, \"t|ket>: A retargetable compiler for NISQ devices,\" Quantum Sci. Technol., vol. 6, no. 1, p. 014003, 2020.",
    "M. Bandic, S. Feld, and C. G. Almudever, \"Full-stack quantum computing systems in the NISQ era: Algorithm-driven and hardware-aware compilation techniques,\" in Proc. DATE, pp. 1-6, 2022.",
    "M. Y. Siraichi, V. F. dos Santos, S. Collange, and F. M. Q. Pereira, \"Qubit allocation,\" in Proc. CGO, pp. 113-125, 2018.",
    "A. Zulehner, A. Paler, and R. Wille, \"An efficient methodology for mapping quantum circuits to the IBM QX architectures,\" IEEE Trans. Comput.-Aided Design Integr. Circuits Syst., vol. 38, no. 7, pp. 1226-1236, 2019.",
    "A. Cowtan, S. Dilkes, R. Duncan, A. Krajenbrink, W. Simmons, and S. Sivarajah, \"On the qubit routing problem,\" in Proc. TQC, LIPIcs, vol. 135, pp. 5:1-5:32, 2019.",
    "A. M. Childs, E. Schoute, and C. M. Unsal, \"Circuit transformations for quantum architectures,\" in Proc. TQC, LIPIcs, vol. 135, pp. 3:1-3:24, 2019.",
    "R. Wille, O. Keszocze, M. Walter, P. Rohrs, A. Chattopadhyay, and R. Drechsler, \"Look-ahead schemes for nearest neighbor optimization of 1D and 2D quantum circuits,\" in Proc. ASP-DAC, pp. 292-297, 2016.",
    "A. Lye, R. Wille, and R. Drechsler, \"Determining the minimal number of SWAP gates for multi-dimensional nearest neighbor quantum circuits,\" in Proc. ASP-DAC, pp. 178-183, 2015.",
    "S. Li, X. Zhou, and Y. Feng, \"Qubit mapping based on subgraph isomorphism and filtered depth-limited search,\" IEEE Trans. Comput., vol. 70, no. 11, pp. 1777-1788, 2021.",
    "T. Itoko, R. Raymond, T. Imamichi, and A. Matsuo, \"Optimization of quantum circuit mapping using gate transformation and commutation,\" Integration, vol. 70, pp. 43-50, 2020.",
    "B. Tan and J. Cong, \"Optimal qubit mapping with simultaneous gate absorption,\" arXiv:2109.06445, 2021.",
    "M. A. Steinberg, S. Feld, C. G. Almudever, M. Marthaler, and J.-M. Reiner, \"Topological-graph dependencies and scaling properties of a heuristic qubit-assignment algorithm,\" IEEE Trans. Quantum Eng., vol. 3, pp. 1-14, 2022.",
    "IBM Quantum, \"SabreLayout API reference,\" IBM Quantum Documentation, 2026.",
    "P. Murali, J. M. Baker, A. Javadi-Abhari, F. T. Chong, and M. Martonosi, \"Noise-adaptive compiler mappings for noisy intermediate-scale quantum computers,\" in Proc. ASPLOS, pp. 1015-1029, 2019.",
    "S. S. Tannu and M. K. Qureshi, \"Not all qubits are created equal: A case for variability-aware policies for NISQ-era quantum computers,\" in Proc. ASPLOS, pp. 987-999, 2019.",
    "S. Nishio, Y. Pan, T. Satoh, H. Amano, and R. Van Meter, \"Extracting success from IBM's 20-qubit machines using error-aware compilation,\" ACM J. Emerg. Technol. Comput. Syst., vol. 16, no. 3, pp. 1-25, 2020.",
    "L. Lao, H. van Someren, I. Ashraf, and C. G. Almudever, \"Timing and resource-aware mapping of quantum circuits to superconducting processors,\" IEEE Trans. Comput.-Aided Design Integr. Circuits Syst., vol. 41, no. 2, pp. 359-371, 2022.",
    "G. G. Guerreschi and J. Park, \"Two-step approach to scheduling quantum circuits,\" Quantum Sci. Technol., vol. 3, no. 4, p. 045003, 2018.",
    "P. Murali et al., \"Full-stack, real-system quantum computer studies: Architectural comparisons and design insights,\" in Proc. ISCA, pp. 527-540, 2019.",
    "S. Hillmich, A. Zulehner, and R. Wille, \"Exploiting quantum teleportation in quantum circuit mapping,\" in Proc. ASP-DAC, pp. 792-797, 2021.",
    "J. M. Baker, C. Duckering, A. Hoover, and F. T. Chong, \"Time-sliced quantum circuit partitioning for modular architectures,\" in Proc. Computing Frontiers, pp. 98-107, 2020.",
    "M. Bandic et al., \"Mapping quantum circuits to modular architectures with QUBO,\" arXiv:2305.06687, 2023.",
    "L. Lao and D. E. Browne, \"2QAN: A quantum compiler for 2-local qubit Hamiltonian simulation algorithms,\" arXiv:2108.02099, 2021.",
    "M. Bandic, H. Zarein, E. Alarcon, and C. G. Almudever, \"On structured design space exploration for mapping of quantum algorithms,\" in Proc. DCIS, pp. 1-6, 2020.",
    "S. Herbert and A. Sengupta, \"Using reinforcement learning to find efficient qubit routing policies for deployment in near-term quantum computers,\" arXiv:1812.11619, 2018.",
    "G. Acampora and R. Schiattarella, \"Deep neural networks for quantum circuit mapping,\" Neural Comput. Appl., vol. 33, pp. 13723-13743, 2021.",
    "T. N. Kipf and M. Welling, \"Semi-supervised classification with graph convolutional networks,\" in Proc. ICLR, 2017.",
    "J. Gilmer, S. S. Schoenholz, P. F. Riley, O. Vinyals, and G. E. Dahl, \"Neural message passing for quantum chemistry,\" in Proc. ICML, PMLR, vol. 70, pp. 1263-1272, 2017.",
    "J. Schulman, F. Wolski, P. Dhariwal, A. Radford, and O. Klimov, \"Proximal policy optimization algorithms,\" arXiv:1707.06347, 2017.",
    "J. Schulman, P. Moritz, S. Levine, M. I. Jordan, and P. Abbeel, \"High-dimensional continuous control using generalized advantage estimation,\" arXiv:1506.02438, 2015.",
    "G. Nannicini, L. S. Bishop, O. Gunluk, and P. Jurcevic, \"Optimal qubit assignment and routing via integer programming,\" ACM Trans. Quantum Comput., vol. 4, no. 1, 2022.",
    "D. Venturelli, M. Do, E. Rieffel, and J. Frank, \"Compiling quantum circuits to realistic hardware architectures using temporal planners,\" Quantum Sci. Technol., vol. 3, no. 2, p. 025004, 2018.",
    "D. Venturelli et al., \"Quantum circuit compilation: An emerging application for automated reasoning,\" in Proc. SPARK, 2019.",
    "M. Bandic, C. G. Almudever, and S. Feld, \"Interaction graph-based characterization of quantum benchmarks for improving quantum circuit mapping techniques,\" Quantum Mach. Intell., vol. 5, p. 40, 2023.",
    "A. W. Cross, L. S. Bishop, S. Sheldon, P. D. Nation, and J. M. Gambetta, \"Validating quantum computers using randomized model circuits,\" Phys. Rev. A, vol. 100, p. 032328, 2019.",
    "R. Blume-Kohout and K. C. Young, \"A volumetric framework for quantum computer benchmarks,\" Quantum, vol. 4, p. 362, 2020.",
    "D. Mills, S. Sivarajah, T. L. Scholten, and R. Duncan, \"Application-motivated, holistic benchmarking of a full quantum computing stack,\" Quantum, vol. 5, p. 415, 2021.",
    "T. Tomesh et al., \"SupermarQ: A scalable quantum benchmark suite,\" in Proc. HPCA, pp. 587-603, 2022.",
    "A. Li and S. Krishnamoorthy, \"QASMBench: A low-level QASM benchmark suite for NISQ evaluation and simulation,\" arXiv:2005.13018, 2020.",
    "E. Farhi, J. Goldstone, and S. Gutmann, \"A quantum approximate optimization algorithm,\" arXiv:1411.4028, 2014.",
    "D. Coppersmith, \"An approximate Fourier transform useful in quantum factoring,\" arXiv:quant-ph/0201067, 2002.",
)

EQUATION_ASSETS = tuple(HERE / "assets" / f"equation_{index}.png" for index in range(1, 5))

RELATED_WORK_LATEX = r"""\section{Related Work}
NISQ compilation sits between abstract circuit descriptions and devices whose
connectivity, native gates, timing, and error characteristics constrain
execution. Standard references establish the computational model and the NISQ
setting, while Qiskit, OpenQL, and t|ket> illustrate retargetable compiler
stacks \citep{preskill2018nisq,bharti2022nisq,nielsenchuang2010,
javadiabhari2024qiskit,khammassi2021openql,sivarajah2020tket,
bandic2022fullstack}. QWeave follows this staged view and reports placement,
routing, lowering, and scheduling separately.

Qubit mapping has been approached through compiler allocation, search, graph
methods, and local rewriting. Siraichi et al. formulate allocation directly;
Zulehner et al. use architecture-aware search; and Cowtan et al. and Childs et
al. study routing independently of a particular device
\citep{siraichi2018allocation,zulehner2019efficient,cowtan2019routing,
childs2019transformations}. Look-ahead nearest-neighbor optimization, exact
minimum-SWAP search, subgraph-isomorphism placement, commutation-aware
transformation, gate absorption, and topology-aware assignment further show
that routing cost depends on both circuit structure and the permitted rewrite
space \citep{wille2016lookahead,lye2015minimal,li2021subgraph,
itoko2020optimization,tan2021gateabsorption,steinberg2022topological}. SABRE's
bidirectional layout and heuristic SWAP search remains a substantive scalable
baseline \citep{li2019sabre,ibmsabre}.

Connectivity is only one hardware constraint. Noise-adaptive, variability-
aware, and error-aware compilers incorporate calibration information, while
timing-aware mappers and scheduling studies distinguish depth from duration
\citep{murali2019noise,tannu2019variability,nishio2020erroraware,
lao2022timing,guerreschi2018twostep}. Full-stack hardware studies and
teleportation-assisted mapping broaden the cost model beyond SWAP count
\citep{murali2019fullstack,hillmich2021teleportation}. QWeave's hardware-cost
study remains synthetic and therefore does not claim calibration-derived
fidelity.

Scaling beyond a monolithic coupling graph introduces partitioning and
communication choices. Time-sliced partitioning, QUBO mapping for modular
architectures, Hamiltonian-simulation-specific compilation, and structured
design-space exploration expose tradeoffs between routing quality, latency,
and compiler runtime \citep{baker2020partitioning,bandic2023qubo,
laobrowne2021twoqan,bandic2020dse}. These methods motivate QWeave's explicit
topology and metric declarations, although the present experiments use only
connected single-device graphs.

Learning-based mapping includes reinforcement-learning policies, learned
initial layouts, and graph-guided search \citep{herbert2018rl,pozzi2020rl,
acampora2021neural,sinha2022gnn}. QWeave uses an edge-conditioned actor--critic
whose encoder draws on graph convolution and message passing, and whose update
uses PPO with generalized advantage estimation
\citep{kipf2017gcn,gilmer2017mpnn,schulman2017ppo,schulman2015gae}. Neither PPO
nor message passing is claimed as novel; the ablations test whether either
helps under the frozen protocol.

Exact optimization and automated planning provide reference solutions for
small instances. Integer programming can jointly model assignment and routing,
whereas temporal planning represents hardware actions and resource constraints
\citep{nannicini2022ip,venturelli2018temporal,venturelli2019reasoning}.
QWeave's CP-SAT component is intentionally smaller: it optimizes only a static
initial-mapping surrogate for at most eight logical qubits and therefore cannot
certify minimum SWAP count or routed depth.

Evaluation methodology also affects conclusions. Interaction-graph
characterization, randomized model circuits, volumetric benchmarks, holistic
stack benchmarks, SupermarQ, and QASMBench emphasize complementary notions of
scale, structure, capability, and application relevance
\citep{bandic2023characterization,cross2019volume,
blumekohout2020volumetric,mills2021holistic,tomesh2022supermarq,
li2020qasmbench}. QWeave consequently freezes circuit families, topologies,
seeds, validity checks, and paired statistics before interpreting method-level
differences.
"""


def _citation(match: re.Match[str]) -> str:
    keys = [key.strip() for key in match.group(2).split(",")]
    numbers = sorted(CITE_NUM[key] for key in keys)
    return ", ".join(f"[{number}]" for number in numbers)


def clean(text: str) -> str:
    text = text.replace(r"\citet{farhi2014qaoa}", r"Farhi et al. \citep{farhi2014qaoa}")
    text = text.replace(r"\citet{coppersmith2002qft}", r"Coppersmith \citep{coppersmith2002qft}")
    text = re.sub(r"\\(citep|citet)\{([^}]+)\}", _citation, text)
    replacements = {
        r"Eq.~\eqref{eq:objective}": "(1)",
        r"Equation~\eqref{eq:route-invariant}": "(2)",
        r"Equation~\eqref{eq:reverse-cx}": "(4)",
        r"Table~\ref{tab:test-results}": "Table I",
        r"Figure~\ref{fig:learning-curves}": "Fig. 1",
        r"\rightarrow": "->",
        r"\leq": "<=",
        r"\ldots": "...",
        r"\mathrm": "",
        r"\texttt": "",
        r"\textsc": "",
        r"\emph": "",
        r"\small": "",
        r"\circ": " o ",
        r"\otimes": " x ",
        r"\{": "{",
        r"\}": "}",
        r"\%": "%",
        r"\_": "_",
        r"~": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\(?:label|ref)\{[^}]+\}", "", text)
    text = re.sub(r"\\[A-Za-z]+", "", text)
    text = text.replace("{", "").replace("}", "").replace("$", "")
    text = text.replace("--", "-")
    return " ".join(text.split())


def _extract() -> tuple[str, str]:
    source = SOURCE.read_text(encoding="utf-8")
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", source, re.S).group(1)
    body = source.split(r"\end{abstract}", 1)[1].split(r"\bibliographystyle", 1)[0]
    body = re.sub(
        r"\\section\{Related Work\}.*?(?=\\section\{Problem Formulation and Deterministic Path\})",
        lambda _match: RELATED_WORK_LATEX + "\n\n",
        body,
        flags=re.S,
    )
    body = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", "\n\n[[RESULTS_TABLE]]\n\n", body, flags=re.S)
    body = re.sub(r"\\begin\{figure\*\}.*?\\end\{figure\*\}", "\n\n[[LEARNING_FIGURE]]\n\n", body, flags=re.S)
    for index in range(4):
        body = re.sub(r"\\begin\{equation\}.*?\\end\{equation\}", f"\n\n[[EQUATION:{index}]]\n\n", body, count=1, flags=re.S)
    return abstract, body


def _roman(number: int) -> str:
    values = ((10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"))
    answer = ""
    for value, symbol in values:
        while number >= value:
            answer += symbol
            number -= value
    return answer


def _set_columns(section, count: int) -> None:
    columns = section._sectPr.xpath("./w:cols")
    element = columns[0] if columns else OxmlElement("w:cols")
    element.set(qn("w:num"), str(count))
    element.set(qn("w:space"), "360")
    if not columns:
        section._sectPr.append(element)


def _configure_section(section, columns: int) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    _set_columns(section, columns)


def _black(run, size: float, bold: bool = False, italic: bool = False, font: str = "Times New Roman") -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def _docx_body(
    document: Document,
    text: str,
    first: bool = False,
    *,
    font_size: float = 10,
    line_spacing: float = 10.8,
) -> None:
    paragraph = document.add_paragraph(clean(text))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0 if first else 0.14)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(line_spacing)
    for run in paragraph.runs:
        _black(run, font_size)


def _docx_section(document: Document, number: int, title: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{_roman(number)}. {clean(title).upper()}")
    _black(run, 10)
    run.font.small_caps = True


def _docx_subsection(document: Document, number: int, title: str) -> None:
    letter = chr(ord("A") + number - 1)
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    _black(p.add_run(f"{letter}. {clean(title)}"), 10, italic=True)


def _docx_unnumbered_section(document: Document, title: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(clean(title).upper())
    _black(run, 10)
    run.font.small_caps = True


def _docx_author_block(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    for index, (name, marker) in enumerate(AUTHORS):
        if index:
            p.add_run(", ")
        if index == 3:
            p.add_run().add_break()
        _black(p.add_run(name), 11)
        marker_run = p.add_run(marker)
        _black(marker_run, 7)
        marker_run.font.superscript = True
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    marker = p.add_run("1")
    _black(marker, 7)
    marker.font.superscript = True
    _black(p.add_run("School of Computer Science and Engineering, VIT Vellore"), 10)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    marker = p.add_run("2")
    _black(marker, 7)
    marker.font.superscript = True
    _black(p.add_run("Faculty Guide, School of Computer Science and Engineering, VIT Vellore"), 10)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    _black(p.add_run("; ".join(EMAILS)), 8.5, font="Courier New")


def _docx_wide_start(document: Document) -> None:
    section = document.add_section(WD_SECTION.CONTINUOUS)
    _configure_section(section, 1)


def _docx_columns_start(document: Document) -> None:
    section = document.add_section(WD_SECTION.CONTINUOUS)
    _configure_section(section, 2)


def _docx_table(document: Document) -> None:
    _docx_wide_start(document)
    caption = document.add_paragraph("TABLE I")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(0)
    _black(caption.runs[0], 8)
    caption = document.add_paragraph("BENCHMARK-V2 TEST SUMMARY. RUNTIME IS MEDIAN WALL TIME IN SECONDS.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(2)
    _black(caption.runs[0], 8)
    rows = (
        ("Method", "Depth", "SWAP", "Runtime", "Fallback"),
        ("Basic", "48.5", "11.5", "0.0022", "0"),
        ("Weighted", "36.0", "0.0", "0.0036", "0"),
        ("SABRE", "46.0", "7.0", "0.0028", "0"),
        ("GNN-PPO", "37.5", "0.0", "0.3473", "4543"),
        ("No-message PPO", "36.0", "0.0", "0.1040", "742"),
        ("Untrained GNN", "198.0", "193.0", "0.6681", "8059"),
    )
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(2.1), Inches(0.8), Inches(0.8), Inches(1.0), Inches(1.0))
    for col, width in zip(table.columns, widths):
        col.width = width
    for row_index, values in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                _black(run, 8, bold=row_index == 0)
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for name in ("top", "insideH", "bottom"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), "000000")
        borders.append(edge)
    for name in ("left", "insideV", "right"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    _docx_columns_start(document)


def _docx_figure(document: Document) -> None:
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)
    _docx_wide_start(document)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(FIGURE), width=Inches(5.9))
    p = document.add_paragraph(
        "Fig. 1. Five-seed learning curves. Lines show the median recent episode return and bands show the interquartile range. Neither curve is monotone, and the GNN curve degrades more strongly late in training."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        _black(run, 8)
    _docx_columns_start(document)


def build_docx(abstract: str, body: str) -> Path:
    document = Document()
    _configure_section(document.sections[0], 1)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(10.8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title.paragraph_format.line_spacing = Pt(27)
    title.paragraph_format.space_after = Pt(6)
    _black(title.add_run(TITLE), 24)
    _docx_author_block(document)
    _docx_columns_start(document)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(9.5)
    _black(p.add_run("Abstract-"), 9, bold=True, italic=True)
    _black(p.add_run(clean(abstract)), 9, italic=True)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    _black(p.add_run("Index Terms-"), 9, bold=True, italic=True)
    _black(p.add_run(KEYWORDS), 9, italic=True)

    tokens = re.split(
        r"(\\section\{[^}]+\}|\\subsection\{[^}]+\}|\\paragraph\{[^}]+\}|\[\[RESULTS_TABLE\]\]|\[\[LEARNING_FIGURE\]\]|\[\[EQUATION:\d\]\])",
        body,
    )
    section_number = 0
    subsection_number = 0
    first = False
    formal_section = False
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        section_match = re.fullmatch(r"\\section\{([^}]+)\}", token)
        subsection_match = re.fullmatch(r"\\subsection\{([^}]+)\}", token)
        runin_match = re.fullmatch(r"\\paragraph\{([^}]+)\}", token)
        equation_match = re.fullmatch(r"\[\[EQUATION:(\d)\]\]", token)
        if section_match:
            section_number += 1
            subsection_number = 0
            section_title = clean(section_match.group(1))
            formal_section = section_title == "Formal Correctness Arguments"
            _docx_section(document, section_number, section_title)
            first = True
        elif subsection_match:
            subsection_number += 1
            _docx_subsection(document, subsection_number, subsection_match.group(1))
            first = True
        elif runin_match:
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            label = clean(runin_match.group(1))
            _black(
                p.add_run(label + (" " if label.endswith((".", ":", ";", "?", "!")) else ". ")),
                11.5 if formal_section else 10,
                bold=True,
                italic=True,
            )
            first = True
        elif token == "[[RESULTS_TABLE]]":
            _docx_table(document)
            first = True
        elif token == "[[LEARNING_FIGURE]]":
            _docx_figure(document)
            first = True
        elif equation_match:
            asset = EQUATION_ASSETS[int(equation_match.group(1))]
            if not asset.exists():
                raise FileNotFoundError(f"missing LaTeX equation render: {asset}")
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(str(asset), width=Inches(3.25))
            first = False
        else:
            for block in re.split(r"\n\s*\n", token):
                if clean(block):
                    _docx_body(
                        document,
                        block,
                        first,
                        font_size=11 if formal_section else 10,
                        line_spacing=12.2 if formal_section else 10.8,
                    )
                    first = False

    _docx_unnumbered_section(document, "References")
    for index, reference in enumerate(REFERENCES, 1):
        p = document.add_paragraph(f"[{index}] {reference}")
        p.paragraph_format.left_indent = Inches(0.17)
        p.paragraph_format.first_line_indent = Inches(-0.17)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(7.5)
        for run in p.runs:
            _black(run, 7.5)

    document.core_properties.title = TITLE
    document.core_properties.author = "; ".join(name for name, _ in AUTHORS)
    document.core_properties.subject = "IEEE conference-style QWeave manuscript"
    document.core_properties.keywords = KEYWORDS
    DOCX_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_OUTPUT)
    return DOCX_OUTPUT


BLACK = colors.HexColor("#000000")


def _register_fonts() -> None:
    fonts = Path("C:/Windows/Fonts")
    pdfmetrics.registerFont(TTFont("TNR", str(fonts / "times.ttf")))
    pdfmetrics.registerFont(TTFont("TNR-Bold", str(fonts / "timesbd.ttf")))
    pdfmetrics.registerFont(TTFont("TNR-Italic", str(fonts / "timesi.ttf")))
    pdfmetrics.registerFont(TTFont("TNR-BoldItalic", str(fonts / "timesbi.ttf")))
    pdfmetrics.registerFont(TTFont("CourierNew", str(fonts / "cour.ttf")))


def _pdf_styles() -> dict[str, ParagraphStyle]:
    return {
        "title": ParagraphStyle("Title", fontName="TNR", fontSize=24, leading=26, alignment=TA_CENTER, textColor=BLACK, spaceAfter=6),
        "authors": ParagraphStyle("Authors", fontName="TNR", fontSize=11, leading=12, alignment=TA_CENTER, textColor=BLACK, spaceAfter=2),
        "affiliation": ParagraphStyle("Affiliation", fontName="TNR", fontSize=10, leading=11, alignment=TA_CENTER, textColor=BLACK),
        "email": ParagraphStyle("Email", fontName="CourierNew", fontSize=8.5, leading=10, alignment=TA_CENTER, textColor=BLACK, spaceAfter=4),
        "abstract": ParagraphStyle("Abstract", fontName="TNR-Italic", fontSize=9, leading=9.8, alignment=TA_JUSTIFY, textColor=BLACK, spaceAfter=2),
        "section": ParagraphStyle("Section", fontName="TNR", fontSize=10, leading=11, alignment=TA_CENTER, textColor=BLACK, spaceBefore=5, spaceAfter=2, keepWithNext=True),
        "subsection": ParagraphStyle("Subsection", fontName="TNR-Italic", fontSize=10, leading=11, alignment=TA_LEFT, textColor=BLACK, spaceBefore=3, spaceAfter=0, keepWithNext=True),
        "runin": ParagraphStyle("Runin", fontName="TNR-BoldItalic", fontSize=10, leading=10.8, alignment=TA_LEFT, textColor=BLACK, spaceBefore=2, keepWithNext=True),
        "formal_runin": ParagraphStyle("FormalRunin", fontName="TNR-BoldItalic", fontSize=11.5, leading=12.5, alignment=TA_LEFT, textColor=BLACK, spaceBefore=3, keepWithNext=True),
        "body": ParagraphStyle("Body", fontName="TNR", fontSize=10, leading=10.8, alignment=TA_JUSTIFY, textColor=BLACK, firstLineIndent=0.14 * inch, spaceAfter=0, hyphenationLang="en_US"),
        "body_first": ParagraphStyle("BodyFirst", fontName="TNR", fontSize=10, leading=10.8, alignment=TA_JUSTIFY, textColor=BLACK, firstLineIndent=0, spaceAfter=0, hyphenationLang="en_US"),
        "formal": ParagraphStyle("Formal", fontName="TNR", fontSize=11, leading=12.2, alignment=TA_JUSTIFY, textColor=BLACK, firstLineIndent=0.14 * inch, spaceAfter=0, hyphenationLang="en_US"),
        "formal_first": ParagraphStyle("FormalFirst", fontName="TNR", fontSize=11, leading=12.2, alignment=TA_JUSTIFY, textColor=BLACK, firstLineIndent=0, spaceAfter=0, hyphenationLang="en_US"),
        "caption": ParagraphStyle("Caption", fontName="TNR", fontSize=8, leading=9, alignment=TA_JUSTIFY, textColor=BLACK, spaceBefore=1, spaceAfter=2),
        "reference": ParagraphStyle("Reference", fontName="TNR", fontSize=8, leading=8.5, alignment=TA_LEFT, textColor=BLACK, leftIndent=0.17 * inch, firstLineIndent=-0.17 * inch),
        "table": ParagraphStyle("Table", fontName="TNR", fontSize=6.6, leading=7.2, alignment=TA_CENTER, textColor=BLACK),
        "table_left": ParagraphStyle("TableLeft", fontName="TNR", fontSize=6.6, leading=7.2, alignment=TA_LEFT, textColor=BLACK),
    }


def _pdf_para(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(clean(text)), style)


def _pdf_equation(index: int, max_width: float):
    asset = EQUATION_ASSETS[index]
    if not asset.exists():
        raise FileNotFoundError(f"missing LaTeX equation render: {asset}")
    with PILImage.open(asset) as raster:
        width, height = raster.size
    target_width = min(max_width, 3.28 * inch)
    image = Image(str(asset), width=target_width, height=target_width * height / width)
    image.hAlign = "CENTER"
    return image


def _pdf_table(styles: dict[str, ParagraphStyle], width: float):
    rows = (
        ("Method", "Depth", "SWAP", "Runtime", "Fallback"),
        ("Basic", "48.5", "11.5", "0.0022", "0"),
        ("Weighted", "36.0", "0.0", "0.0036", "0"),
        ("SABRE", "46.0", "7.0", "0.0028", "0"),
        ("GNN-PPO", "37.5", "0.0", "0.3473", "4543"),
        ("No-message PPO", "36.0", "0.0", "0.1040", "742"),
        ("Untrained GNN", "198.0", "193.0", "0.6681", "8059"),
    )
    caption = Paragraph("TABLE I<br/>BENCHMARK-V2 TEST SUMMARY. RUNTIME IS MEDIAN WALL TIME IN SECONDS.", styles["caption"])
    data = []
    for row_index, row in enumerate(rows):
        data.append([
            Paragraph(("<b>" + escape(value) + "</b>") if row_index == 0 else escape(value), styles["table_left"] if col == 0 else styles["table"])
            for col, value in enumerate(row)
        ])
    table = Table(data, colWidths=(0.34 * width, 0.145 * width, 0.145 * width, 0.19 * width, 0.18 * width), repeatRows=1)
    table.setStyle(TableStyle([
        ("TEXTCOLOR", (0, 0), (-1, -1), BLACK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ("LEFTPADDING", (0, 0), (-1, -1), 0.5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0.5),
        ("LINEABOVE", (0, 0), (-1, 0), 0.6, BLACK),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, BLACK),
        ("LINEBELOW", (0, -1), (-1, -1), 0.6, BLACK),
    ]))
    return KeepTogether([caption, table])


def build_pdf(abstract: str, body: str) -> Path:
    _register_fonts()
    styles = _pdf_styles()
    page_width, page_height = LETTER
    left_margin = right_margin = 0.625 * inch
    top_margin = 0.75 * inch
    bottom_margin = 1.0 * inch
    gap = 0.25 * inch
    usable_width = page_width - left_margin - right_margin
    usable_height = page_height - top_margin - bottom_margin
    column_width = (usable_width - gap) / 2
    title_height = 1.92 * inch
    title_frame = Frame(left_margin, page_height - top_margin - title_height, usable_width, title_height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="title")
    first_height = usable_height - title_height - 0.08 * inch
    first_left = Frame(left_margin, bottom_margin, column_width, first_height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="first-left")
    first_right = Frame(left_margin + column_width + gap, bottom_margin, column_width, first_height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="first-right")
    left = Frame(left_margin, bottom_margin, column_width, usable_height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="left")
    right = Frame(left_margin + column_width + gap, bottom_margin, column_width, usable_height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="right")
    PDF_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = BaseDocTemplate(str(PDF_OUTPUT), pagesize=LETTER, title=TITLE, author="; ".join(name for name, _ in AUTHORS), subject="IEEE conference-style QWeave manuscript", keywords=KEYWORDS)
    document.addPageTemplates([
        PageTemplate(id="First", frames=[title_frame, first_left, first_right], autoNextPageTemplate="TwoColumn"),
        PageTemplate(id="TwoColumn", frames=[left, right]),
    ])
    author_html = ", ".join(f"{escape(name)}<super>{marker}</super>" for name, marker in AUTHORS[:3]) + "<br/>" + ", ".join(f"{escape(name)}<super>{marker}</super>" for name, marker in AUTHORS[3:])
    story = [
        Paragraph(escape(TITLE), styles["title"]),
        Paragraph(author_html, styles["authors"]),
        Paragraph("<super>1</super>School of Computer Science and Engineering, VIT Vellore", styles["affiliation"]),
        Paragraph("<super>2</super>Faculty Guide, School of Computer Science and Engineering, VIT Vellore", styles["affiliation"]),
        Paragraph("; ".join(EMAILS), styles["email"]),
        FrameBreak(),
        Paragraph("<b>Abstract-</b>" + escape(clean(abstract)), styles["abstract"]),
        Paragraph("<b>Index Terms-</b>" + escape(KEYWORDS), styles["abstract"]),
    ]
    tokens = re.split(
        r"(\\section\{[^}]+\}|\\subsection\{[^}]+\}|\\paragraph\{[^}]+\}|\[\[RESULTS_TABLE\]\]|\[\[LEARNING_FIGURE\]\]|\[\[EQUATION:\d\]\])",
        body,
    )
    section_number = 0
    subsection_number = 0
    first = False
    formal_section = False
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        section_match = re.fullmatch(r"\\section\{([^}]+)\}", token)
        subsection_match = re.fullmatch(r"\\subsection\{([^}]+)\}", token)
        runin_match = re.fullmatch(r"\\paragraph\{([^}]+)\}", token)
        equation_match = re.fullmatch(r"\[\[EQUATION:(\d)\]\]", token)
        if section_match:
            section_number += 1
            subsection_number = 0
            section_title = clean(section_match.group(1))
            formal_section = section_title == "Formal Correctness Arguments"
            story.append(Paragraph(f"{_roman(section_number)}. {escape(section_title.upper())}", styles["section"]))
            first = True
        elif subsection_match:
            subsection_number += 1
            story.append(Paragraph(f"{chr(64 + subsection_number)}. {escape(clean(subsection_match.group(1)))}", styles["subsection"]))
            first = True
        elif runin_match:
            label = clean(runin_match.group(1))
            story.append(Paragraph(escape(label + ("" if label.endswith((".", ":", ";", "?", "!")) else ".")), styles["formal_runin"] if formal_section else styles["runin"]))
            first = True
        elif token == "[[RESULTS_TABLE]]":
            story.append(_pdf_table(styles, column_width))
            first = True
        elif token == "[[LEARNING_FIGURE]]":
            if not FIGURE.exists():
                raise FileNotFoundError(FIGURE)
            image = Image(str(FIGURE), width=column_width, height=column_width * 0.38)
            image.hAlign = "CENTER"
            story.append(KeepTogether([image, Paragraph("Fig. 1. Five-seed learning curves. Lines show the median recent episode return and bands show the interquartile range. Neither curve is monotone, and the GNN curve degrades more strongly late in training.", styles["caption"])]))
            first = True
        elif equation_match:
            story.append(_pdf_equation(int(equation_match.group(1)), column_width))
            first = False
        else:
            for block in re.split(r"\n\s*\n", token):
                if clean(block):
                    if formal_section:
                        paragraph_style = styles["formal_first"] if first else styles["formal"]
                    else:
                        paragraph_style = styles["body_first"] if first else styles["body"]
                    story.append(_pdf_para(block, paragraph_style))
                    first = False
    story.append(Paragraph("REFERENCES", styles["section"]))
    for index, reference in enumerate(REFERENCES, 1):
        story.append(Paragraph(f"[{index}] {escape(reference)}", styles["reference"]))
    document.build(story)
    return PDF_OUTPUT


def build_tex() -> Path:
    source = SOURCE.read_text(encoding="utf-8")
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", source, re.S).group(1).strip()
    body = source.split(r"\end{abstract}", 1)[1].split(r"\bibliographystyle", 1)[0].strip()
    body = re.sub(
        r"\\section\{Related Work\}.*?(?=\\section\{Problem Formulation and Deterministic Path\})",
        lambda _match: RELATED_WORK_LATEX + "\n\n",
        body,
        flags=re.S,
    )
    body = body.replace(r"\citet{farhi2014qaoa}", r"Farhi et al.~\cite{farhi2014qaoa}")
    body = body.replace(r"\citet{coppersmith2002qft}", r"Coppersmith~\cite{coppersmith2002qft}")
    body = body.replace(r"\citep", r"\cite")
    body = body.replace(
        r"\section{Formal Correctness Arguments}",
        r"\section{Formal Correctness Arguments}" "\n" r"\begingroup\fontsize{11}{12.5}\selectfont",
    )
    body = body.replace(
        r"\section{Learned Router}",
        r"\endgroup" "\n" r"\section{Learned Router}",
    )
    body = re.sub(
        r"(\\begin\{equation\}.*?\\end\{equation\})",
        lambda match: "{\\large\n" + match.group(1) + "\n}",
        body,
        flags=re.S,
    )
    body = body.replace("../results/benchmark_v2/figures/learning_curves.png", "../../../results/benchmark_v2/figures/learning_curves.png")
    preamble = r"""\documentclass[conference]{IEEEtran}
\usepackage{amsmath,amssymb}
\usepackage{booktabs}
\usepackage{graphicx}
\usepackage{cite}
\usepackage{url}
\title{QWeave: Correctness-First Deterministic and Learned Qubit Routing}
\author{\IEEEauthorblockN{Atharva Sheersh Pandey$^{1}$, Diptesh Das$^{1}$, Shrivardhini N$^{1}$, Haridasu Sreedhar$^{1}$, Prof. Bhuvaneswari M$^{2}$}
\IEEEauthorblockA{$^{1}$School of Computer Science and Engineering, VIT Vellore\\
$^{2}$Faculty Guide, School of Computer Science and Engineering, VIT Vellore\\
\{atharva.sheersh2024, diptesh.das2024, shrivardhini.n2024, haridasu.sreedhar2024\}@vitstudent.ac.in}}
\begin{document}
\maketitle
\begin{abstract}
"""
    bibitems = "\n".join(
        "\\bibitem{" + key + "} " + reference.replace("&", r"\&")
        for key, reference in zip(REFERENCE_KEYS, REFERENCES)
    )
    ending = "\n\\begin{thebibliography}{00}\n" + bibitems + "\n\\end{thebibliography}\n\\end{document}\n"
    tex = preamble + abstract + "\n\\end{abstract}\n\\begin{IEEEkeywords}\n" + KEYWORDS + "\n\\end{IEEEkeywords}\n\n" + body + "\n" + ending
    TEX_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    TEX_OUTPUT.write_text(tex, encoding="utf-8")
    return TEX_OUTPUT


def build_all() -> tuple[Path, Path, Path]:
    abstract, body = _extract()
    tex = build_tex()
    docx = build_docx(abstract, body)
    pdf = build_pdf(abstract, body)
    return tex, docx, pdf


if __name__ == "__main__":
    for artifact in build_all():
        print(artifact)
